# -*- coding: utf-8 -*-
"""M1.5 导入器测试：对照真实表结构、空表、缺列、多列、合成历史数据。

用法：python tools/test_xlsx_import.py [真实表路径]
"""
import io
import json
import os
import sys
import tempfile

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import xlsx_to_json as X  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)

# 真实表路径：不写死在代码里（那会把本机路径与用户名带进开源仓库）。
# 优先命令行参数 → 环境变量 → 仓库上一级目录里的同名文件。
# 找不到就跳过这组用例（开源用户没有这张表也必须能跑通），而不是失败。
def resolve_real():
    if len(sys.argv) > 1:
        return sys.argv[1]
    env = os.environ.get("GL_REAL_XLSX")
    if env:
        return env
    guess = os.path.join(os.path.dirname(ROOT), "半年计划打卡表_26周.xlsx")
    return guess if os.path.exists(guess) else None


REAL = resolve_real()

failed = 0


def check(name, ok, extra=""):
    global failed
    print(("PASS  " if ok else "FAIL  ") + name + ("  " + str(extra) if extra else ""))
    if not ok:
        failed += 1


def tmpfile(suffix=".xlsx"):
    fd, p = tempfile.mkstemp(suffix=suffix)
    os.close(fd)
    return p


# ---------- 1. 真实表 ----------
print("— 1 真实表结构对齐 —")
if REAL and os.path.exists(REAL):
    d = X.convert(REAL)
    check("schema 正确", d["schema"] == "growth-loop.import.v1")
    check("工作表全部识别", len(d["warnings"]) == 0, d["warnings"])
    check("26 周记录完整", len(d["weekly"]) == 26, len(d["weekly"]))
    check("周记录带起止日期", d["weekly"][0]["monday"] == "2026-10-19" and d["weekly"][25]["sunday"] == "2027-04-18",
          d["weekly"][0]["monday"] + " ~ " + d["weekly"][25]["sunday"])
    check("里程碑 14 条", len(d["milestones"]) == 14, len(d["milestones"]))
    check("资金配置 7 条", len(d["allocation"]) == 7, len(d["allocation"]))
    check("书单 31 条", len(d["books"]) == 31, len(d["books"]))
    check("量化素材 5 条", len(d["assets"]) == 5, len(d["assets"]))
    check("基线抓到内容", len(d["baseline"]) >= 3, len(d["baseline"]))
    check("目标行解析出体脂区间", d["metricTargets"].get("bodyfat", "").find("15") >= 0, d["metricTargets"].get("bodyfat"))
    check("目标行解析出腰围", "82" in str(d["metricTargets"].get("waist", "")), d["metricTargets"].get("waist"))
    check("指标月份为空（真实表未填）", d["metrics"] == {}, list(d["metrics"].keys()))
    check("工作表名原样保留", len(d["source"]["sheets"]) == 8, len(d["source"]["sheets"]))
    check("真实表里没有多余的 week 字段", set(d["weekly"][0].keys()) <= {"week", "monday", "sunday", "stage", "aiTask", "financeTask", "reading", "hours", "workout", "calorieDays", "listenDays", "speakCount", "aiHours", "reviewed", "note"}, list(d["weekly"][0].keys()))
else:
    print("SKIP  跳过真实表用例（没有这张表；用 GL_REAL_XLSX 或命令行参数指定路径）")


# ---------- 2. 空表不崩 ----------
print("\n— 2 空表 —")
p = tmpfile()
wb = openpyxl.Workbook()
wb.active.title = "空白"
wb.save(p)
d = X.convert(p)
check("空表不抛异常", True)
check("结构都在且为空", d["weekly"] == [] and d["metrics"] == {} and d["milestones"] == [])
check("缺工作表有警告", any("缺少工作表" in w for w in d["warnings"]), d["warnings"][:1])
os.remove(p)


# ---------- 3. 缺列 ----------
print("\n— 3 缺列 —")
p = tmpfile()
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "指标追踪(12个月)"
ws.append(["月份", "体重kg"])          # 只有两列
ws.append(["M1", 70])
ws.append(["M2", 69])
ws2 = wb.create_sheet("周打卡(26周)")
ws2.append(["周次"])      # 只有周次列
ws2.append([3])           # 一行数据，其余列都不存在
wb.save(p)
d = X.convert(p)
check("缺列仍能取到有值的项", d["metrics"].get("M1", {}).get("weight") == 70, d["metrics"])
check("缺的指标不出现在结果里", "bodyfat" not in d["metrics"].get("M1", {}))
check("周表缺列仍能读出周次", d["weekly"] == [{"week": 3}], d["weekly"])
os.remove(p)


# ---------- 4. 多列 / 列序变化 ----------
print("\n— 4 多列与列序变化 —")
p = tmpfile()
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "指标追踪(12个月)"
ws.append(["备注A", "体脂%", "月份", "体重kg", "备注B"])   # 列序乱、夹带无关列
ws.append(["x", "25", "M1", "70", "y"])
ws.append(["x", None, "M2", None, "y"])                     # 空行不该产出
wb.save(p)
d = X.convert(p)
check("列序变化不影响取值", d["metrics"].get("M1") == {"bodyfat": 25, "weight": 70}, d["metrics"])
check("全空行不产出", "M2" not in d["metrics"], list(d["metrics"].keys()))
os.remove(p)


# ---------- 5. 合成历史数据：导入后计算正确 ----------
print("\n— 5 合成历史数据 —")
sample = X.make_sample()
check("合成 12 个月指标", len(sample["metrics"]) == 12, len(sample["metrics"]))
check("合成 26 周记录", len(sample["weekly"]) == 26, len(sample["weekly"]))
months = sorted(sample["metrics"].keys())
check("月份连续", months[0] == "2026-10" and months[-1] == "2027-09", months[0] + " ~ " + months[-1])
w = sample["metrics"]["2026-10"]
check("体脂逐月下降", sample["metrics"]["2027-09"]["bodyfat"] < w["bodyfat"],
      f'{w["bodyfat"]} → {sample["metrics"]["2027-09"]["bodyfat"]}')
total_hours = sum(r["hours"] for r in sample["weekly"])
check("周投入合计可算", 400 < total_hours < 800, round(total_hours, 1))

# 测试不许改仓库文件（否则 git status 变脏）。写到临时目录，再与已提交的样例逐字节比对。
tmp = tmpfile(".json")
with io.open(tmp, "w", encoding="utf-8", newline="\n") as f:
    json.dump(sample, f, ensure_ascii=False, indent=2)
committed = os.path.join(ROOT, "data", "sample.json")
if os.path.exists(committed):
    same = open(committed, "rb").read() == open(tmp, "rb").read()
    check("data/sample.json 与生成器逐字节一致", same,
          "" if same else "重跑：python tools/xlsx_to_json.py --sample -o data/sample.json")
    raw = open(committed, "rb").read()
    check("样例文件是 LF 行尾", raw.count(b"\r\n") == 0, "CRLF=%d" % raw.count(b"\r\n"))
else:
    check("data/sample.json 存在", False, "跑：python tools/xlsx_to_json.py --sample -o data/sample.json")
os.remove(tmp)

print("\n全部通过" if failed == 0 else "\n失败 %d 项" % failed)
sys.exit(0 if failed == 0 else 1)
